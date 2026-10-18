# File: apps/dashboard/views.py
"""
Views cho trang quản trị riêng (Dashboard) của Admin:
Thống kê doanh thu, đơn hàng, khách hàng, hành vi AI và truy xuất nguồn gốc.
Sử dụng staff_member_required để phân quyền quản trị viên / nhân viên.
"""
import io
import json
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from django.shortcuts import get_object_or_404, redirect
from django.contrib import messages
from django.views.generic import TemplateView
from django.views import View
from django.http import HttpResponse
from django.contrib.admin.views.decorators import staff_member_required
from django.utils.decorators import method_decorator
from django.db.models import Sum, Count, Q
from django.db import models
from datetime import timedelta
from django.utils import timezone
from apps.accounts.models import User
from apps.products.models import Product, Category, Supplier
from apps.cart_orders.models import Order, OrderItem
from apps.behavior.models import UserEvent
from apps.traceability.models import ProductBatch


@method_decorator(staff_member_required, name='dispatch')
class DashboardHomeView(TemplateView):
    """
    Trang tổng quan Dashboard: Doanh thu 30 ngày, biểu đồ xu hướng Chart.js,
    phân nhóm khách hàng AI và top nông sản bán chạy.
    """
    template_name = 'dashboard/index.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        d30 = timezone.now() - timedelta(days=30)
        orders_30d = Order.objects.filter(created_at__gte=d30).exclude(status='cancelled')
        ctx['total_revenue'] = orders_30d.aggregate(s=Sum('final_amount'))['s'] or 0
        ctx['total_orders'] = orders_30d.count()
        ctx['total_customers'] = User.objects.filter(is_staff=False).count()
        ctx['total_products'] = Product.objects.filter(is_active=True).count()
        ctx['top_products'] = (
            OrderItem.objects.filter(order__created_at__gte=d30)
            .values('product_name').annotate(total=Sum('quantity')).order_by('-total')[:10]
        )

        # 1. Dữ liệu biểu đồ doanh thu 30 ngày (sử dụng Python timezone để tương thích tuyệt đối MySQL trên Windows)
        today = timezone.now().date()
        daily_revenue_map = {today - timedelta(days=i): 0.0 for i in range(29, -1, -1)}
        recent_orders = (
            Order.objects.filter(created_at__gte=timezone.now() - timedelta(days=30))
            .exclude(status='cancelled')
            .values('created_at', 'final_amount')
        )
        for ord_val in recent_orders:
            dt_local = timezone.localtime(ord_val['created_at']).date()
            if dt_local in daily_revenue_map:
                daily_revenue_map[dt_local] += float(ord_val['final_amount'] or 0)

        chart_revenue_labels = [d.strftime('%d/%m') for d in daily_revenue_map.keys()]
        chart_revenue_data = list(daily_revenue_map.values())

        # 2. Phân nhóm khách hàng AI cho biểu đồ tròn (Doughnut Chart)
        segment_labels_map = {
            'new': 'Khách mới',
            'interested': 'Quan tâm',
            'loyal': 'Thường xuyên',
            'at_risk': 'Nguy cơ rời bỏ',
        }
        segments_raw = list(User.objects.filter(is_staff=False).values('segment').annotate(count=Count('id')))
        ctx['segments'] = segments_raw

        chart_segment_labels = []
        chart_segment_data = []
        for s in segments_raw:
            seg_code = s['segment'] or 'other'
            chart_segment_labels.append(segment_labels_map.get(seg_code, 'Khác' if seg_code == 'other' else seg_code))
            chart_segment_data.append(s['count'])

        # 3. Top 5 nông sản bán chạy cho biểu đồ cột (Bar Chart)
        top_items = list(ctx['top_products'][:5])
        chart_top_labels = [item['product_name'] for item in top_items]
        chart_top_data = [int(item['total'] or 0) for item in top_items]

        # Truyền dạng chuỗi JSON an toàn cho JavaScript
        ctx['chart_revenue_labels'] = json.dumps(chart_revenue_labels)
        ctx['chart_revenue_data'] = json.dumps(chart_revenue_data)
        ctx['chart_segment_labels'] = json.dumps(chart_segment_labels)
        ctx['chart_segment_data'] = json.dumps(chart_segment_data)
        ctx['chart_top_labels'] = json.dumps(chart_top_labels)
        ctx['chart_top_data'] = json.dumps(chart_top_data)

        return ctx


@method_decorator(staff_member_required, name='dispatch')
class DashboardProductsView(TemplateView):
    """
    Quản lý danh sách nông sản và tồn kho với KPI chi tiết và bộ lọc thông minh.
    """
    template_name = 'dashboard/products.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        all_prods = Product.objects.select_related('category', 'supplier').all()

        # Thống kê KPI nông sản & kho
        ctx['total_products'] = all_prods.count()
        ctx['active_products'] = all_prods.filter(is_active=True).count()
        ctx['low_stock_products'] = all_prods.filter(stock__gt=0, stock__lt=10).count()
        ctx['out_of_stock_products'] = all_prods.filter(stock__lte=0).count()

        # Tổng giá trị nông sản trong kho
        total_inv_val = 0
        for p in all_prods:
            if p.stock and p.stock > 0 and p.price:
                total_inv_val += (p.price * p.stock)
        ctx['total_inventory_value'] = total_inv_val

        # Danh mục cho bộ lọc nhanh
        ctx['categories'] = Category.objects.all().order_by('name')

        # Xử lý tìm kiếm và bộ lọc
        qs = all_prods.order_by('-created_at')
        cat_id = self.request.GET.get('category')
        stock_filter = self.request.GET.get('stock')
        search_query = self.request.GET.get('q', '').strip()

        if cat_id and cat_id.isdigit():
            qs = qs.filter(category_id=int(cat_id))

        if stock_filter == 'active':
            qs = qs.filter(is_active=True)
        elif stock_filter == 'low':
            qs = qs.filter(stock__gt=0, stock__lt=10)
        elif stock_filter == 'out':
            qs = qs.filter(stock__lte=0)
        elif stock_filter == 'inactive':
            qs = qs.filter(is_active=False)

        if search_query:
            qs = qs.filter(
                Q(name__icontains=search_query) |
                Q(origin__icontains=search_query) |
                Q(supplier__name__icontains=search_query)
            )

        ctx['current_category'] = cat_id or ''
        ctx['current_stock_filter'] = stock_filter or 'all'
        ctx['search_query'] = search_query
        ctx['products'] = qs
        return ctx


@method_decorator(staff_member_required, name='dispatch')
class DashboardOrdersView(TemplateView):
    """
    Quản lý và theo dõi đơn hàng nông sản với số liệu thống kê chi tiết.
    """
    template_name = 'dashboard/orders.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        all_orders = Order.objects.all()

        # Thống kê số lượng theo trạng thái
        ctx['total_orders_count'] = all_orders.count()
        ctx['pending_count'] = all_orders.filter(status='pending').count()
        ctx['confirmed_count'] = all_orders.filter(status='confirmed').count()
        ctx['shipping_count'] = all_orders.filter(status='shipping').count()
        ctx['delivered_count'] = all_orders.filter(status='delivered').count()
        ctx['cancelled_count'] = all_orders.filter(status='cancelled').count()

        # Đơn VietQR chưa có ảnh bill
        ctx['vietqr_pending_bill_count'] = all_orders.filter(
            payment_method='vietqr', status='pending'
        ).filter(models.Q(payment_proof='') | models.Q(payment_proof__isnull=True)).count()

        # Doanh thu các đơn giao thành công
        ctx['total_delivered_revenue'] = (
            all_orders.filter(status='delivered').aggregate(s=Sum('final_amount'))['s'] or 0
        )

        # Lấy danh sách 100 đơn hàng mới nhất kèm các mặt hàng (items)
        orders_qs = (
            Order.objects.select_related('user')
            .prefetch_related('items__product')
            .order_by('-created_at')
        )

        status_param = self.request.GET.get('status')
        if status_param and status_param in dict(Order.STATUS_CHOICES):
            orders_qs = orders_qs.filter(status=status_param)
        elif status_param == 'need_bill':
            orders_qs = orders_qs.filter(
                payment_method='vietqr', status='pending'
            ).filter(models.Q(payment_proof='') | models.Q(payment_proof__isnull=True))

        ctx['current_status'] = status_param or 'all'
        ctx['orders'] = orders_qs[:100]
        return ctx


@method_decorator(staff_member_required, name='dispatch')
class DashboardUpdateOrderStatusView(View):
    """
    Cập nhật trạng thái đơn hàng từ Dashboard.
    Hỗ trợ xử lý thông minh:
    - Đơn COD: Cập nhật trạng thái ngay lập tức.
    - Đơn VietQR: Cho phép Admin duyệt ngay nếu đã đối soát tài khoản ngân hàng thực tế,
      hoặc đính kèm file bill thay khách hàng; hệ thống tự động sinh chứng từ đối soát nội bộ hợp lệ.
    - Hủy đơn 'pending': Tự động hoàn lại số lượng tồn kho sản phẩm.
    """
    def post(self, request, order_id):
        order = get_object_or_404(Order, id=order_id)
        target_status = request.POST.get('status')
        if not target_status:
            messages.error(request, "Thiếu trạng thái cần cập nhật.")
            return redirect('dashboard:orders')

        # Xử lý hủy đơn: hoàn lại tồn kho nếu đơn đang chờ duyệt
        if target_status == 'cancelled' and order.status == 'pending':
            from apps.cart_orders.services import cancel_order
            cancel_order(order)
            messages.success(request, f"✅ Đã hủy đơn hàng #{order.order_code} và hoàn trả số lượng tồn kho thành công!")
            return redirect('dashboard:orders')

        # Xử lý đơn VietQR khi xác nhận hoặc chuyển tiếp vận chuyển
        if target_status in ('confirmed', 'shipping', 'delivered') and order.payment_method == 'vietqr':
            # Trường hợp 1: Admin tải trực tiếp file bill đính kèm
            if 'payment_proof' in request.FILES:
                order.payment_proof = request.FILES['payment_proof']
                order.payment_proof_uploaded_at = timezone.now()
                admin_name = request.user.username if request.user.is_authenticated else 'Admin'
                note_line = f"[Admin {admin_name} đã tải lên bill đối soát lúc {timezone.now().strftime('%d/%m/%Y %H:%M')}]"
                order.note = (order.note + "\n" + note_line).strip() if order.note else note_line

            # Trường hợp 2: Chưa có bill -> kiểm tra cờ xác nhận đối soát từ Admin
            elif not order.payment_proof:
                admin_verified = request.POST.get('admin_verified') in ('1', 'true', 'True')
                if admin_verified:
                    from apps.cart_orders.services import generate_admin_verification_bill
                    admin_name = request.user.username if request.user.is_authenticated else 'Admin'
                    proof_file = generate_admin_verification_bill(order.order_code, admin_username=admin_name)
                    order.payment_proof.save(f"admin_verified_{order.order_code}.png", proof_file, save=False)
                    order.payment_proof_uploaded_at = timezone.now()
                    note_line = f"[Admin {admin_name} đã đối soát ngân hàng thành công lúc {timezone.now().strftime('%d/%m/%Y %H:%M')}]"
                    order.note = (order.note + "\n" + note_line).strip() if order.note else note_line
                else:
                    messages.error(
                        request,
                        f"❌ KHÔNG THỂ XÁC NHẬN: Đơn hàng #{order.order_code} thanh toán qua VietQR nhưng KHÁCH HÀNG CHƯA TẢI LÊN BILL CHUYỂN KHOẢN! Vui lòng chờ khách tải bill đối soát."
                    )
                    return redirect('dashboard:orders')


        order.status = target_status
        order.save()
        messages.success(request, f"✅ Đã cập nhật đơn hàng #{order.order_code} sang trạng thái: {order.get_status_display()} thành công!")
        return redirect('dashboard:orders')



@method_decorator(staff_member_required, name='dispatch')
class DashboardCustomersView(TemplateView):
    """
    Quản lý danh sách người dùng và phân loại nhóm hành vi AI.
    """
    template_name = 'dashboard/customers.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        base_customers = User.objects.filter(is_staff=False)

        # Thống kê phân nhóm AI
        ctx['total_customers'] = base_customers.count()
        ctx['loyal_count'] = base_customers.filter(segment='loyal').count()
        ctx['interested_count'] = base_customers.filter(segment='interested').count()
        ctx['new_count'] = base_customers.filter(segment='new').count()
        ctx['at_risk_count'] = base_customers.filter(segment='at_risk').count()
        ctx['normal_count'] = base_customers.filter(segment='normal').count()

        # Danh sách kèm số đơn hàng đã đặt
        qs = base_customers.annotate(
            total_orders=Count('orders')
        ).order_by('-date_joined')

        segment_param = self.request.GET.get('segment')
        search_query = self.request.GET.get('q', '').strip()

        if segment_param and segment_param in dict(User.SEGMENT_CHOICES):
            qs = qs.filter(segment=segment_param)

        if search_query:
            qs = qs.filter(
                Q(username__icontains=search_query) |
                Q(email__icontains=search_query) |
                Q(phone__icontains=search_query) |
                Q(first_name__icontains=search_query) |
                Q(last_name__icontains=search_query)
            )

        ctx['current_segment'] = segment_param or 'all'
        ctx['search_query'] = search_query
        ctx['customers'] = qs
        return ctx


@method_decorator(staff_member_required, name='dispatch')
class DashboardBehaviorView(TemplateView):
    """
    Thống kê tần suất sự kiện hành vi trong 30 ngày và danh sách 100 sự kiện mới nhất.
    """
    template_name = 'dashboard/behavior.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        d30 = timezone.now() - timedelta(days=30)
        events_30d = UserEvent.objects.filter(created_at__gte=d30)

        # Thống kê 6 loại hành vi tương tác chính
        ctx['total_events_30d'] = events_30d.count()
        ctx['view_count'] = events_30d.filter(event_type='view').count()
        ctx['search_count'] = events_30d.filter(event_type='search').count()
        ctx['cart_count'] = events_30d.filter(event_type='add_to_cart').count()
        ctx['purchase_count'] = events_30d.filter(event_type='purchase').count()
        ctx['review_count'] = events_30d.filter(event_type='review').count()
        ctx['chat_count'] = events_30d.filter(event_type='chat').count()

        ctx['event_counts'] = (
            events_30d.values('event_type').annotate(count=Count('id')).order_by('-count')
        )

        # Danh sách sự kiện kèm bộ lọc
        qs = (
            UserEvent.objects.select_related('user', 'product', 'category')
            .order_by('-created_at')
        )
        event_filter = self.request.GET.get('type')
        search_query = self.request.GET.get('q', '').strip()

        if event_filter and event_filter in dict(UserEvent.EVENT_TYPES):
            qs = qs.filter(event_type=event_filter)

        if search_query:
            qs = qs.filter(
                Q(keyword__icontains=search_query) |
                Q(user__username__icontains=search_query) |
                Q(product__name__icontains=search_query) |
                Q(session_key__icontains=search_query)
            )

        ctx['current_type'] = event_filter or 'all'
        ctx['search_query'] = search_query
        ctx['recent_events'] = qs[:100]
        return ctx


@method_decorator(staff_member_required, name='dispatch')
class DashboardTraceabilityView(TemplateView):
    """
    Quản lý lô hàng nông sản và danh sách mã QR truy xuất.
    """
    template_name = 'dashboard/traceability.html'

    def get_context_data(self, **kwargs):
        ctx = super().get_context_data(**kwargs)
        all_batches = ProductBatch.objects.select_related('product__category', 'product__supplier').all()

        # Thống kê KPI lô hàng
        ctx['total_batches'] = all_batches.count()
        ctx['has_qr_count'] = all_batches.exclude(qr_code='').count()
        ctx['certified_count'] = all_batches.exclude(certification='').count()
        ctx['products_with_batch'] = Product.objects.filter(batches__isnull=False).distinct().count()

        # Bộ lọc và tìm kiếm
        qs = all_batches.order_by('-created_at')
        search_query = self.request.GET.get('q', '').strip()
        cert_filter = self.request.GET.get('cert')

        if cert_filter:
            qs = qs.filter(certification__icontains=cert_filter)

        if search_query:
            qs = qs.filter(
                Q(batch_code__icontains=search_query) |
                Q(product__name__icontains=search_query) |
                Q(planting_location__icontains=search_query)
            )

        ctx['search_query'] = search_query
        ctx['current_cert'] = cert_filter or 'all'
        ctx['batches'] = qs
        return ctx


@method_decorator(staff_member_required, name='dispatch')
class ExportOrdersExcelView(View):
    """
    Xuất báo cáo toàn bộ danh sách đơn hàng ra file Excel (.xlsx) chuẩn chỉnh.
    Được định dạng chuyên nghiệp với Header màu thương hiệu, viền ô và căn lề tự động.
    """
    def get(self, request, *args, **kwargs):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Báo Cáo Đơn Hàng"

        # Định dạng Header
        header_font = Font(name="Segoe UI", size=11, bold=True, color="FFFFFF")
        header_fill = PatternFill(start_color="10B981", end_color="10B981", fill_type="solid")
        header_align = Alignment(horizontal="center", vertical="center", wrap_text=True)

        headers = [
            "STT", "Mã đơn hàng", "Ngày đặt", "Khách hàng", "SĐT nhận hàng",
            "Địa chỉ giao hàng", "Tiền hàng (đ)", "Phí vận chuyển (đ)", "Tổng thanh toán (đ)",
            "Trạng thái", "Phương thức", "Ghi chú"
        ]
        ws.append(headers)

        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = header_align

        status_map = dict(Order.STATUS_CHOICES)
        pay_map = dict(Order.PAYMENT_CHOICES)

        orders = Order.objects.all().order_by('-created_at')
        thin_border = Border(
            left=Side(style='thin', color='E2E8F0'),
            right=Side(style='thin', color='E2E8F0'),
            top=Side(style='thin', color='E2E8F0'),
            bottom=Side(style='thin', color='E2E8F0')
        )

        for idx, ord in enumerate(orders, 1):
            created_str = timezone.localtime(ord.created_at).strftime('%d/%m/%Y %H:%M') if ord.created_at else ''
            row = [
                idx,
                ord.order_code,
                created_str,
                ord.full_name or (ord.user.username if ord.user else 'Khách'),
                ord.phone,
                ord.address,
                float(ord.total_amount),
                float(ord.shipping_fee),
                float(ord.final_amount),
                status_map.get(ord.status, ord.status),
                pay_map.get(ord.payment_method, ord.payment_method),
                ord.note or ''
            ]
            ws.append(row)
            row_idx = idx + 1
            for col_idx in range(1, len(row) + 1):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if col_idx in [7, 8, 9]:
                    cell.number_format = '#,##0'
                    cell.alignment = Alignment(horizontal='right')
                elif col_idx in [1, 3, 5, 10, 11]:
                    cell.alignment = Alignment(horizontal='center')

        # Tự động căn chỉnh độ rộng cột
        for col in ws.columns:
            max_len = max(len(str(cell.value or '')) for cell in col)
            col_letter = openpyxl.utils.get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

        buffer = io.BytesIO()
        wb.save(buffer)
        buffer.seek(0)

        filename = f"AgriSense_Orders_{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
        response = HttpResponse(
            buffer.getvalue(),
            content_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
        )
        response['Content-Disposition'] = f'attachment; filename="{filename}"'
        return response
