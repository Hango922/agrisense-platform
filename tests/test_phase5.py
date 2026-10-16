# File: tests/test_phase5.py
"""
Kiểm thử toàn diện Dashboard Admin Phase 5:
1. Phân quyền: User thường bị chặn, Staff/Admin truy cập được
2. Số liệu thống kê Dashboard (doanh thu, đơn hàng, khách hàng, sản phẩm, phân nhóm)
3. Các trang phân hệ con của Dashboard
"""
from django.test import TestCase, Client
from apps.accounts.models import User
from apps.products.models import Product, Category, Supplier
from apps.cart_orders.models import Order


class Phase5ComprehensiveTests(TestCase):
    def setUp(self):
        self.cat = Category.objects.create(name='Rau củ', slug='rau-cu')
        self.sup = Supplier.objects.create(name='Đà Lạt Xanh', address='Đà Lạt')
        self.prod = Product.objects.create(
            name='Cà rốt', slug='ca-rot', category=self.cat, supplier=self.sup,
            price=25000, stock=50, unit='kg'
        )
        self.normal_user = User.objects.create_user(username='normal_user', password='Password@123')
        self.admin_user = User.objects.create_superuser(
            username='admin_user',
            email='admin@test.com',
            password='AdminPassword@123'
        )
        self.order = Order.objects.create(
            order_code='ORD-TEST-DASHBOARD',
            user=self.normal_user,
            full_name='Khách Mẫu Dashboard',
            phone='0901234567',
            address='TP. Hồ Chí Minh',
            total_amount=220000,
            shipping_fee=30000,
            final_amount=250000,
            status='confirmed',
            payment_method='cod'
        )
        self.client = Client()

    def test_permission_denied_for_normal_user(self):
        self.client.force_login(self.normal_user)
        response = self.client.get('/dashboard/')
        self.assertIn(response.status_code, [302, 403])

    def test_admin_dashboard_metrics(self):
        self.client.force_login(self.admin_user)
        response = self.client.get('/dashboard/')
        self.assertEqual(response.status_code, 200)
        ctx = response.context
        self.assertGreaterEqual(ctx['total_revenue'], 250000)
        self.assertGreaterEqual(ctx['total_orders'], 1)
        self.assertGreaterEqual(ctx['total_customers'], 1)
        self.assertGreaterEqual(ctx['total_products'], 1)

    def test_dashboard_subpages(self):
        self.client.force_login(self.admin_user)
        subpages = [
            '/dashboard/products/',
            '/dashboard/orders/',
            '/dashboard/customers/',
            '/dashboard/behavior/',
            '/dashboard/traceability/',
        ]
        for url in subpages:
            res = self.client.get(url)
            self.assertEqual(res.status_code, 200, f"Page {url} returned {res.status_code}")

    def test_admin_dashboard_chartjs_data(self):
        """Kiểm tra context chứa đầy đủ dữ liệu JSON cho 3 biểu đồ Chart.js và HTML chứa canvas."""
        import json
        self.client.force_login(self.admin_user)
        response = self.client.get('/dashboard/')
        self.assertEqual(response.status_code, 200)
        ctx = response.context

        # Kiểm tra 3 bộ dữ liệu biểu đồ
        self.assertIn('chart_revenue_labels', ctx)
        self.assertIn('chart_revenue_data', ctx)
        self.assertIn('chart_segment_labels', ctx)
        self.assertIn('chart_top_labels', ctx)

        rev_labels = json.loads(ctx['chart_revenue_labels'])
        rev_data = json.loads(ctx['chart_revenue_data'])
        self.assertEqual(len(rev_labels), 30)
        self.assertEqual(len(rev_data), 30)

        # Kiểm tra thẻ canvas và CDN trong template HTML
        self.assertContains(response, '<canvas id="revenueChart"></canvas>')
        self.assertContains(response, '<canvas id="segmentChart"></canvas>')
        self.assertContains(response, '<canvas id="topProductsChart"></canvas>')
        self.assertContains(response, 'chart.js')

    def test_export_orders_excel_permission(self):
        """Người dùng thường không được phép tải file Excel báo cáo đơn hàng."""
        self.client.force_login(self.normal_user)
        response = self.client.get('/dashboard/orders/export-excel/')
        self.assertIn(response.status_code, [302, 403])

    def test_export_orders_excel_success(self):
        """Quản trị viên xuất file Excel báo cáo đơn hàng định dạng chuẩn .xlsx thành công."""
        import io
        import openpyxl
        self.client.force_login(self.admin_user)
        response = self.client.get('/dashboard/orders/export-excel/')
        self.assertEqual(response.status_code, 200)
        self.assertEqual(
            response['Content-Type'],
            'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
        )
        self.assertIn('.xlsx', response['Content-Disposition'])

        # Đọc và xác thực cấu trúc workbook Excel
        wb = openpyxl.load_workbook(io.BytesIO(response.content))
        ws = wb.active
        self.assertEqual(ws.title, "Báo Cáo Đơn Hàng")

        # Kiểm tra dòng Header
        headers = [cell.value for cell in ws[1]]
        self.assertIn("Mã đơn hàng", headers)
        self.assertIn("Tổng thanh toán (đ)", headers)
        self.assertIn("Trạng thái", headers)

        # Kiểm tra dữ liệu đơn hàng mẫu có trong file
        order_codes = [ws.cell(row=r, column=2).value for r in range(2, ws.max_row + 1)]
        self.assertIn(self.order.order_code, order_codes)

    def test_dashboard_update_vietqr_order_blocked_without_proof(self):
        """Admin không thể xác nhận đơn hàng VietQR nếu khách chưa tải ảnh bill."""
        vietqr_order = Order.objects.create(
            order_code='ORD-TEST-VQ-NOBILL',
            user=self.normal_user,
            full_name='Khách Test VietQR',
            phone='0988777666',
            address='Hà Nội',
            total_amount=100000,
            final_amount=130000,
            status='pending',
            payment_method='vietqr'
        )
        self.client.force_login(self.admin_user)
        response = self.client.post(
            f'/dashboard/orders/{vietqr_order.id}/update-status/',
            {'status': 'confirmed'},
            follow=True
        )
        self.assertEqual(response.status_code, 200)
        vietqr_order.refresh_from_db()
        # Vẫn ở pending vì chưa có bill
        self.assertEqual(vietqr_order.status, 'pending')
        self.assertContains(response, 'KHÔNG THỂ XÁC NHẬN')

    def test_dashboard_update_vietqr_order_succeeds_with_proof(self):
        """Admin xác nhận thành công đơn hàng VietQR khi khách đã tải bill chuyển khoản."""
        from django.core.files.uploadedfile import SimpleUploadedFile
        fake_bill = SimpleUploadedFile("bill.png", b"\x89PNG\r\n\x1a\ncontent", content_type="image/png")
        vietqr_order = Order.objects.create(
            order_code='ORD-TEST-VQ-BILL',
            user=self.normal_user,
            full_name='Khách Test VietQR Đã Có Bill',
            phone='0988777666',
            address='Hà Nội',
            total_amount=100000,
            final_amount=130000,
            status='pending',
            payment_method='vietqr',
            payment_proof=fake_bill
        )
        self.client.force_login(self.admin_user)
        response = self.client.post(
            f'/dashboard/orders/{vietqr_order.id}/update-status/',
            {'status': 'confirmed'},
            follow=True
        )
        self.assertEqual(response.status_code, 200)
        vietqr_order.refresh_from_db()
        self.assertEqual(vietqr_order.status, 'confirmed')

        if vietqr_order.payment_proof:
            vietqr_order.payment_proof.delete(save=False)

